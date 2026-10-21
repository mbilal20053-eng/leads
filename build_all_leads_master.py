import os
import sys
import re
import csv
import time
import duckdb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from leadgen.sources import _overture_release
from leadgen.verticals._common import is_social
from leadgen.quality import is_junk_name, normalize_phone

# Regional markets across the United States
US_REGIONS = [
    {"market": "Austin & Central TX", "state": "TX", "south": 30.1, "west": -98.0, "north": 30.6, "east": -97.5},
    {"market": "Los Angeles & Orange County, CA", "state": "CA", "south": 33.6, "west": -118.6, "north": 34.3, "east": -117.8},
    {"market": "New York Metro & Tri-State", "state": "NY", "south": 40.5, "west": -74.2, "north": 41.0, "east": -73.7},
    {"market": "Chicago & Cook County, IL", "state": "IL", "south": 41.6, "west": -88.0, "north": 42.1, "east": -87.5},
    {"market": "Miami & South Florida", "state": "FL", "south": 25.6, "west": -80.4, "north": 26.3, "east": -80.0},
    {"market": "Dallas - Fort Worth, TX", "state": "TX", "south": 32.5, "west": -97.4, "north": 33.1, "east": -96.6},
    {"market": "Houston Metro, TX", "state": "TX", "south": 29.6, "west": -95.6, "north": 30.1, "east": -95.1},
    {"market": "Atlanta Metro, GA", "state": "GA", "south": 33.6, "west": -84.6, "north": 34.0, "east": -84.2},
    {"market": "Phoenix Metro, AZ", "state": "AZ", "south": 33.3, "west": -112.3, "north": 33.7, "east": -111.9},
    {"market": "Seattle & King County, WA", "state": "WA", "south": 47.4, "west": -122.4, "north": 47.8, "east": -122.2},
    {"market": "Denver & Front Range, CO", "state": "CO", "south": 39.6, "west": -105.1, "north": 40.0, "east": -104.8},
    {"market": "San Francisco Bay Area, CA", "state": "CA", "south": 37.3, "west": -122.4, "north": 37.9, "east": -122.0},
]

# Industry verticals to classify and target
VERTICAL_DEFS = {
    "Home Services & Trades": {
        "keywords": ["plumb", "electric", "roof", "hvac", "contractor", "landscap", "clean", "paint", "pest", "handyman", "locksmith", "remodel"],
        "pitch": "No dedicated website — homeowners searching on Google can't find them; pitch a fast quote/booking site that captures local search traffic.",
    },
    "Restaurants & Food": {
        "keywords": ["restaurant", "cafe", "baker", "pizza", "bar", "food", "coffee", "taco", "diner", "grill", "bistro", "catering"],
        "pitch": "No owned website — missing direct online orders and reservations; pitch a mobile-friendly menu & online ordering site.",
    },
    "Beauty & Salons": {
        "keywords": ["salon", "barber", "spa", "nail", "hair", "beauty", "esthetic", "tattoo", "lash", "skincare", "massage"],
        "pitch": "No website or relying on social — pitch a 24/7 client booking calendar and service portfolio site.",
    },
    "Automotive Services": {
        "keywords": ["auto", "mechanic", "tire", "car_wash", "detail", "towing", "collision", "brake", "transmission"],
        "pitch": "No website — drivers in need of repairs go to competitors on Google Maps; pitch an instant service quote and emergency contact site.",
    },
    "Healthcare & Dental": {
        "keywords": ["dentist", "doctor", "physician", "chiro", "clinic", "health", "optomet", "therap", "dental", "pediatric"],
        "pitch": "No official web portal — prospective patients need insurance info and appointment booking; pitch a HIPAA-compliant practice website.",
    },
    "Professional Services & Real Estate": {
        "keywords": ["legal", "lawyer", "attorney", "account", "tax", "consult", "real_estate", "realtor", "insurance", "notary", "advisor"],
        "pitch": "Missing professional web presence — high-ticket clients look for trust and credentials; pitch an authoritative firm website.",
    },
    "Pet Services": {
        "keywords": ["veterinar", "pet", "dog", "groom", "animal", "vet", "cat", "kennel"],
        "pitch": "No web presence — pet owners search online for local care; pitch an online booking and boarding reservation website.",
    },
    "Fitness & Wellness": {
        "keywords": ["fitness", "gym", "yoga", "martial_arts", "pilates", "crossfit", "trainer", "studio", "boxing"],
        "pitch": "No owned website — pitch a class schedule and automated membership signup page.",
    },
    "Retail & Boutiques": {
        "keywords": ["store", "shop", "boutique", "retail", "apparel", "clothing", "florist", "jewel", "gift"],
        "pitch": "No ecommerce or catalog site — pitch a local store showcase with Google Shopping integration.",
    },
}

EXCLUDE_CATEGORIES = {
    "elementary_school", "middle_school", "high_school", "school", "preschool",
    "place_of_worship", "church", "baptist_place_of_worship", "catholic_church",
    "cemetery", "police_station", "fire_station", "courthouse", "city_hall",
    "government_office", "embassy", "military_base", "mountain_bike_trail",
    "park", "playground"
}

def detect_vertical(category: str, name: str) -> tuple[str, str]:
    text = f"{category or ''} {name or ''}".lower()
    for v_name, v_info in VERTICAL_DEFS.items():
        if any(k in text for k in v_info["keywords"]):
            return v_name, v_info["pitch"]
    return "General Local Business", "No owned website — pitch a professional business website to rank above local competitors."

def clean_phone(p: str) -> str:
    if not p:
        return ""
    norm = normalize_phone(p)
    return norm if norm else p

def format_category(cat: str) -> str:
    if not cat:
        return "Local Business"
    return cat.replace("_", " ").title()

def main():
    print("=" * 66)
    print("  Generating Master All-Leads Dataset Across the United States")
    print("=" * 66)

    con = duckdb.connect()
    con.execute("INSTALL httpfs; LOAD httpfs; SET s3_region='us-west-2';")
    release = _overture_release(con)
    print(f"Overture Release: {release}")

    s3 = f"s3://overturemaps-us-west-2/release/{release}/theme=places/type=place/*"

    all_leads = []
    seen = set()

    for idx, reg in enumerate(US_REGIONS, 1):
        m_name = reg["market"]
        print(f"[{idx}/{len(US_REGIONS)}] Gathering leads in {m_name}...")

        # 1. Query businesses with NO website (Urgent Need)
        sql_no_site = f"""
        SELECT names.primary AS name,
               taxonomy.primary AS category,
               websites[1] AS website,
               CASE WHEN length(phones)>0 THEN phones[1] END AS phone,
               addresses[1].freeform AS address,
               addresses[1].locality AS city,
               addresses[1].region AS state,
               addresses[1].postcode AS zip,
               addresses[1].country AS country
        FROM read_parquet('{s3}', hive_partitioning=1)
        WHERE bbox.xmin BETWEEN {reg['west']} AND {reg['east']}
          AND bbox.ymin BETWEEN {reg['south']} AND {reg['north']}
          AND names.primary IS NOT NULL
          AND (websites IS NULL OR length(websites) = 0)
          AND length(phones) > 0
          AND (addresses[1].country = 'US' OR addresses[1].country IS NULL)
        LIMIT 65;
        """

        # 2. Query businesses with ONLY a social page
        sql_social = f"""
        SELECT names.primary AS name,
               taxonomy.primary AS category,
               websites[1] AS website,
               CASE WHEN length(phones)>0 THEN phones[1] END AS phone,
               addresses[1].freeform AS address,
               addresses[1].locality AS city,
               addresses[1].region AS state,
               addresses[1].postcode AS zip,
               addresses[1].country AS country
        FROM read_parquet('{s3}', hive_partitioning=1)
        WHERE bbox.xmin BETWEEN {reg['west']} AND {reg['east']}
          AND bbox.ymin BETWEEN {reg['south']} AND {reg['north']}
          AND names.primary IS NOT NULL
          AND length(websites) > 0
          AND (
              lower(websites[1]) LIKE '%facebook.com%' OR
              lower(websites[1]) LIKE '%instagram.com%' OR
              lower(websites[1]) LIKE '%linktr.ee%' OR
              lower(websites[1]) LIKE '%tiktok.com%' OR
              lower(websites[1]) LIKE '%twitter.com%' OR
              lower(websites[1]) LIKE '%x.com%'
          )
          AND (addresses[1].country = 'US' OR addresses[1].country IS NULL)
        LIMIT 35;
        """

        try:
            rows_no_site = con.execute(sql_no_site).fetchall()
        except Exception as e:
            print(f"  Warning on no-site query: {e}")
            rows_no_site = []

        try:
            rows_social = con.execute(sql_social).fetchall()
        except Exception as e:
            print(f"  Warning on social query: {e}")
            rows_social = []

        region_added = 0

        # Process NO WEBSITE leads
        for r in rows_no_site:
            name, cat, site, phone, addr, city, state, zipcode, country = r
            if is_junk_name(name) or cat in EXCLUDE_CATEGORIES or not phone:
                continue

            dedupe_key = (name.strip().lower(), phone.strip())
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)

            vertical, default_pitch = detect_vertical(cat, name)
            clean_p = clean_phone(phone)
            clean_st = state or reg["state"]
            clean_ci = city or m_name.split("&")[0].split("-")[0].strip()

            lead = {
                "Lead ID": len(all_leads) + 1,
                "Industry / Vertical": vertical,
                "Opportunity Type": "NO WEBSITE (Urgent Need)",
                "Tier": "A",
                "Score": 80 if addr else 75,
                "Business Name": name.strip(),
                "Category": format_category(cat),
                "Phone": clean_p,
                "Online Presence": "None (No Website Found)",
                "Address": addr or "",
                "City": clean_ci,
                "State": clean_st,
                "Zip": zipcode or "",
                "Country": "United States",
                "Why It's a Lead": "Active commercial phone listed but ZERO company website exists",
                "Cold Pitch / Suggested Opener": default_pitch
            }
            all_leads.append(lead)
            region_added += 1

        # Process SOCIAL-ONLY leads
        for r in rows_social:
            name, cat, site, phone, addr, city, state, zipcode, country = r
            if is_junk_name(name) or cat in EXCLUDE_CATEGORIES or not site or not is_social(site):
                continue

            dedupe_key = (name.strip().lower(), (phone or "").strip())
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)

            vertical, _ = detect_vertical(cat, name)
            clean_p = clean_phone(phone) if phone else "Unlisted (Contact via Social)"
            clean_st = state or reg["state"]
            clean_ci = city or m_name.split("&")[0].split("-")[0].strip()
            clean_s = site.split("?")[0].strip()

            lead = {
                "Lead ID": len(all_leads) + 1,
                "Industry / Vertical": vertical,
                "Opportunity Type": "SOCIAL-ONLY (Needs Owned Site)",
                "Tier": "A",
                "Score": 75 if phone else 70,
                "Business Name": name.strip(),
                "Category": format_category(cat),
                "Phone": clean_p,
                "Online Presence": clean_s,
                "Address": addr or "",
                "City": clean_ci,
                "State": clean_st,
                "Zip": zipcode or "",
                "Country": "United States",
                "Why It's a Lead": "Renting audience on social platform; no owned domain or website",
                "Cold Pitch / Suggested Opener": (
                    f"Relying entirely on {clean_s.split('/')[2]} — pitch an owned web portal "
                    f"that captures Google search traffic and protects from platform algorithm changes."
                )
            }
            all_leads.append(lead)
            region_added += 1

        print(f"  -> Added {region_added} leads (Total Collected: {len(all_leads)})")

    # Sort all leads by Vertical, then Score desc, then State, then Business Name
    all_leads.sort(key=lambda x: (x["Industry / Vertical"], -x["Score"], x["State"], x["Business Name"]))
    # Re-assign sequential IDs
    for idx, lead in enumerate(all_leads, 1):
        lead["Lead ID"] = idx

    print(f"\n============================================================")
    print(f"  Total Verified US Leads Collected: {len(all_leads)}")
    print(f"============================================================")

    # ──────────────── 1. Export Master CSV for Google Sheets ────────────────
    csv_file = "all_leads_master_google_sheets.csv"
    fieldnames = list(all_leads[0].keys()) if all_leads else []
    with open(csv_file, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_leads)
    print(f"[OK] Master Google Sheets CSV created: {csv_file}")

    # ──────────────── 2. Export Master Multi-Tab Excel Workbook ────────────────
    xlsx_file = "all_leads_master_google_sheets.xlsx"
    wb = openpyxl.Workbook()

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    tier_fill = PatternFill(start_color="E8F8EC", end_color="E8F8EC", fill_type="solid")
    border_thin = Side(border_style="thin", color="DDE2E8")
    cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

    def write_sheet(ws, rows_data, title):
        ws.title = title
        # Headers
        for col_idx, h in enumerate(fieldnames, 1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = cell_border

        # Rows
        for row_idx, lead in enumerate(rows_data, 2):
            for col_idx, h in enumerate(fieldnames, 1):
                val = lead[h]
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = Font(name="Segoe UI", size=10)
                c.fill = tier_fill
                c.border = cell_border
                if h in ("Lead ID", "Tier", "Score", "State", "Zip", "Country"):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)
        ws.freeze_panes = "A2"

    # Tab 1: All Leads Combined
    ws_master = wb.active
    write_sheet(ws_master, all_leads, "All Leads Master")

    # Tabs for each industry vertical
    vertical_groups = {}
    for l in all_leads:
        v = l["Industry / Vertical"]
        vertical_groups.setdefault(v, []).append(l)

    for v_name, v_leads in vertical_groups.items():
        # Excel sheet name max 31 chars
        tab_name = v_name.replace("&", "and")[:30]
        ws_sub = wb.create_sheet(title=tab_name)
        write_sheet(ws_sub, v_leads, tab_name)

    wb.save(xlsx_file)
    print(f"[OK] Master Multi-Tab Excel created: {xlsx_file}")

    # Copy to GUI output folder as well
    os.makedirs("gui/_output", exist_ok=True)
    import shutil
    shutil.copyfile(csv_file, "gui/_output/all_leads_master_google_sheets.csv")
    shutil.copyfile(xlsx_file, "gui/_output/all_leads_master_google_sheets.xlsx")
    print(f"[OK] Copied to gui/_output/ for web download.")

if __name__ == "__main__":
    main()
